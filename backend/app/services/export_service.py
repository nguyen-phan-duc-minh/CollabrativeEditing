"""
Export Service - Convert document content to various formats
"""
from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import markdown2
import html2text
from bs4 import BeautifulSoup
from io import BytesIO
import re
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_LEFT, TA_CENTER


class ExportService:
    """Service for exporting documents to various formats"""
    
    @staticmethod
    def html_to_docx(html_content: str, title: str = "Document") -> BytesIO:
        """
        Convert HTML content to DOCX format
        
        Args:
            html_content: HTML string from TipTap editor
            title: Document title
            
        Returns:
            BytesIO buffer containing DOCX data
        """
        doc = DocxDocument()
        
        # Add title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Parse HTML
        soup = BeautifulSoup(html_content, 'html.parser')
        
        # Process each HTML element
        for element in soup.descendants:
            if element.name == 'p':
                text = element.get_text()
                if text.strip():
                    para = doc.add_paragraph(text)
                    
                    # Apply formatting
                    for run in para.runs:
                        # Check for strong/bold
                        if element.find('strong'):
                            run.bold = True
                        # Check for em/italic
                        if element.find('em'):
                            run.italic = True
                        # Check for u/underline
                        if element.find('u'):
                            run.underline = True
                            
            elif element.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
                level = int(element.name[1])
                text = element.get_text()
                if text.strip():
                    doc.add_heading(text, level=level)
                    
            elif element.name == 'ul':
                for li in element.find_all('li', recursive=False):
                    text = li.get_text()
                    if text.strip():
                        doc.add_paragraph(text, style='List Bullet')
                        
            elif element.name == 'ol':
                for li in element.find_all('li', recursive=False):
                    text = li.get_text()
                    if text.strip():
                        doc.add_paragraph(text, style='List Number')
                        
            elif element.name == 'blockquote':
                text = element.get_text()
                if text.strip():
                    para = doc.add_paragraph(text)
                    para.style = 'Quote'
                    
            elif element.name == 'code':
                text = element.get_text()
                if text.strip():
                    para = doc.add_paragraph(text)
                    for run in para.runs:
                        run.font.name = 'Courier New'
                        run.font.size = Pt(10)
        
        # Save to buffer
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return buffer
    
    @staticmethod
    def html_to_pdf(html_content: str, title: str = "Document") -> BytesIO:
        """
        Convert HTML content to PDF format using ReportLab
        
        Args:
            html_content: HTML string from TipTap editor
            title: Document title
            
        Returns:
            BytesIO buffer containing PDF data
        """
        buffer = BytesIO()
        
        # Create PDF document
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=inch,
            leftMargin=inch,
            topMargin=inch,
            bottomMargin=inch,
            title=title
        )
        
        # Create styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            spaceAfter=30,
            alignment=TA_CENTER
        )
        
        # Parse HTML and convert to PDF elements
        story = []
        
        # Add title
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 24))
        
        # Parse HTML content
        soup = BeautifulSoup(html_content, 'html.parser')
        
        # Convert HTML tags to PDF paragraphs
        for element in soup.descendants:
            if element.name == 'p':
                text = element.get_text()
                if text.strip():
                    # Escape special characters for ReportLab
                    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    story.append(Paragraph(text, styles['Normal']))
                    story.append(Spacer(1, 12))
                    
            elif element.name in ['h1', 'h2', 'h3', 'h4']:
                text = element.get_text()
                if text.strip():
                    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    heading_level = element.name[1]
                    heading_style = styles.get(f'Heading{heading_level}', styles['Heading1'])
                    story.append(Paragraph(text, heading_style))
                    story.append(Spacer(1, 12))
                    
            elif element.name == 'ul':
                for li in element.find_all('li', recursive=False):
                    text = li.get_text()
                    if text.strip():
                        text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                        story.append(Paragraph('• ' + text, styles['Normal']))
                        story.append(Spacer(1, 6))
                        
            elif element.name == 'ol':
                for idx, li in enumerate(element.find_all('li', recursive=False), 1):
                    text = li.get_text()
                    if text.strip():
                        text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                        story.append(Paragraph(f'{idx}. ' + text, styles['Normal']))
                        story.append(Spacer(1, 6))
                        
            elif element.name == 'blockquote':
                text = element.get_text()
                if text.strip():
                    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    quote_style = ParagraphStyle(
                        'Quote',
                        parent=styles['Normal'],
                        leftIndent=36,
                        rightIndent=36,
                        textColor=RGBColor(0.4, 0.4, 0.4)
                    )
                    story.append(Paragraph(text, quote_style))
                    story.append(Spacer(1, 12))
        
        # Build PDF
        doc.build(story)
        buffer.seek(0)
        
        return buffer
    
    @staticmethod
    def html_to_markdown(html_content: str) -> str:
        """
        Convert HTML content to Markdown format
        
        Args:
            html_content: HTML string from TipTap editor
            
        Returns:
            Markdown string
        """
        h = html2text.HTML2Text()
        h.ignore_links = False
        h.body_width = 0  # Don't wrap lines
        h.unicode_snob = True
        h.skip_internal_links = False
        
        markdown = h.handle(html_content)
        
        return markdown.strip()
    
    @staticmethod
    def html_to_text(html_content: str) -> str:
        """
        Convert HTML content to plain text
        
        Args:
            html_content: HTML string from TipTap editor
            
        Returns:
            Plain text string
        """
        soup = BeautifulSoup(html_content, 'html.parser')
        text = soup.get_text(separator='\n')
        
        # Clean up extra whitespace
        lines = [line.strip() for line in text.split('\n')]
        text = '\n'.join(line for line in lines if line)
        
        return text
